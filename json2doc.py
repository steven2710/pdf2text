
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn


# =========================
# 1. Cấu hình file đầu vào/đầu ra
# =========================

INPUT_FILE = "extracted/07-MIG User Guide.json"
OUTPUT_FILE = "extracted/07-MIG User Guide.docx"


# =========================
# 2. Loại bỏ ký tự không hợp lệ trong XML 1.0
# =========================

def clean_xml_text(text):
    """
    Loại bỏ các ký tự không hợp lệ đối với XML 1.0.
    Giữ lại tab, newline, carriage return và các ký tự Unicode hợp lệ.
    """
    if not isinstance(text, str):
        text = str(text)

    cleaned = []

    for ch in text:
        code = ord(ch)

        # XML 1.0 cho phép tab, newline, carriage return
        if code in (0x09, 0x0A, 0x0D):
            cleaned.append(ch)

        # Các ký tự hợp lệ trong vùng BMP
        elif 0x20 <= code <= 0xD7FF:
            cleaned.append(ch)

        elif 0xE000 <= code <= 0xFFFD:
            cleaned.append(ch)

        # Các ký tự Unicode ngoài BMP, trừ FFFE và FFFF
        elif 0x10000 <= code <= 0x10FFFF:
            if (code & 0xFFFF) not in (0xFFFE, 0xFFFF):
                cleaned.append(ch)

    return "".join(cleaned)


# =========================
# 3. Đọc dữ liệu JSON
# =========================

raw = Path(INPUT_FILE).read_text(encoding="utf-8")

try:
    pages = json.loads(raw)

except json.JSONDecodeError:
    # Trường hợp file có thêm nội dung bên ngoài mảng JSON
    start = raw.find("[")
    end = raw.rfind("]")

    if start == -1 or end == -1:
        raise ValueError(
            "Không tìm thấy mảng JSON trong file đầu vào."
        )

    pages = json.loads(raw[start:end + 1])


# Sắp xếp các trang theo số trang nguồn
pages = sorted(pages, key=lambda p: int(p["page"]))


# =========================
# 4. Khởi tạo tài liệu Word
# =========================

doc = Document()

section = doc.sections[0]
section.top_margin = Inches(0.6)
section.bottom_margin = Inches(0.6)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)


# Thiết lập font mặc định
style = doc.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(9)
style._element.get_or_add_rPr().get_or_add_rFonts().set(
    qn("w:eastAsia"), "Arial"
)


# =========================
# 5. Chuyển nội dung từng trang sang Word
# =========================

for index, page in enumerate(pages):

    page_number = page["page"]

    # Làm sạch nội dung trước khi đưa vào Word
    text = clean_xml_text(page.get("text", ""))

    # Tiêu đề đánh dấu trang nguồn
    heading = doc.add_paragraph()
    heading.paragraph_format.space_after = Pt(4)

    run = heading.add_run(f"[Trang nguồn: {page_number}]")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "Arial"
    )

    # Thêm từng dòng văn bản của trang
    for line in text.splitlines():

        # Làm sạch thêm một lần để bảo đảm an toàn
        line = clean_xml_text(line)

        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0

        run = paragraph.add_run(line)
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Arial"
        )

    # Ngắt trang, trừ trang cuối cùng
    if index < len(pages) - 1:
        doc.add_page_break()


# =========================
# 6. Lưu file Word
# =========================

output_path = Path(OUTPUT_FILE)

# Tạo thư mục đầu ra nếu chưa tồn tại
output_path.parent.mkdir(parents=True, exist_ok=True)

doc.save(output_path)

print(f"Đã chuyển {len(pages)} trang nguồn sang: {output_path}")
